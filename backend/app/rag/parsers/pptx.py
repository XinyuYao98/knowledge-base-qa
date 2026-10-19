"""PPTX 解析器（.pptx）。

## 核心决策：备注（notes_slide）必须提取

演示文稿有个特点：**页面上只有关键词，真正的解释在备注里。**

用户问「Q3 目标是多少」，页面上只有「经营目标」四个字，
答案 8600 万在备注里。不读备注，这份 PPT 就等于没内容。

所以这里把正文与备注合并成一个块，并标注来源。
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.rag.parsers.base import (
    Block,
    BlockType,
    ParsedDoc,
    ParseError,
    ParseErrorCode,
    TableCell,
    WarnCode,
)
from app.rag.parsers.registry import BaseParser


class PptxParser(BaseParser):
    extensions = ("pptx",)
    format_name = "pptx"
    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def _parse(self, path: Path, doc: ParsedDoc) -> None:
        try:
            prs = Presentation(str(path))
        except Exception as e:
            raise ParseError(
                ParseErrorCode.CORRUPT_FILE,
                f"无法打开 PPT 文件：{type(e).__name__}",
                detail={"error": str(e)[:200]},
            ) from e

        slide_count = len(prs.slides)
        doc.pages = slide_count
        doc.meta["slide_count"] = slide_count

        if slide_count == 0:
            raise ParseError(ParseErrorCode.EMPTY_FILE, "演示文稿没有幻灯片")

        table_counter = 0
        notes_chars = 0
        body_chars = 0
        picture_count = 0

        for idx, slide in enumerate(prs.slides, start=1):
            title = self._slide_title(slide)
            body_parts: list[str] = []
            tables_here: list[Block] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    txt = shape.text_frame.text.strip()
                    if txt:
                        body_parts.append(txt)
                elif shape.has_table:
                    table_counter += 1
                    blk = self._table_block(shape.table, table_counter, idx)
                    if blk is not None:
                        tables_here.append(blk)
                elif shape.shape_type == 13:  # PICTURE
                    picture_count += 1

            body_text = "\n".join(body_parts)
            notes_text = self._notes(slide)

            body_chars += len(body_text)
            notes_chars += len(notes_text)

            # 正文与备注合并：这是本解析器的核心
            combined = body_text
            if notes_text:
                combined = f"{body_text}\n\n【备注】{notes_text}" if body_text else f"【备注】{notes_text}"

            if combined.strip():
                doc.blocks.append(
                    Block(
                        type=BlockType.PARAGRAPH,
                        text=combined.strip(),
                        page=idx,
                        extra={
                            "slide": idx,
                            "title": title,
                            "has_notes": bool(notes_text),
                            "body_chars": len(body_text),
                            "notes_chars": len(notes_text),
                        },
                    )
                )

            doc.blocks.extend(tables_here)
            doc.page_char_counts.append(len(combined))

        if picture_count:
            doc.warn(
                WarnCode.IMAGE_SKIPPED,
                f"演示文稿含 {picture_count} 张图片，图片内文字未提取（不支持 OCR）",
                count=picture_count,
            )

        if notes_chars > body_chars * 2 and notes_chars > 200:
            doc.warn(
                WarnCode.GENERIC,
                f"备注内容（{notes_chars} 字）远多于页面正文（{body_chars} 字），"
                f"已合并提取，检索时可命中备注内容",
                body_chars=body_chars,
                notes_chars=notes_chars,
            )

        if not doc.blocks:
            raise ParseError(ParseErrorCode.EMPTY_FILE, "演示文稿中没有可提取的文字")

        doc.meta["table_count"] = table_counter
        doc.meta["body_chars"] = body_chars
        doc.meta["notes_chars"] = notes_chars

    # ------------------------------------------------------------------
    @staticmethod
    def _slide_title(slide) -> str:  # type: ignore[no-untyped-def]
        try:
            if slide.shapes.title is not None:
                return slide.shapes.title.text.strip()
        except Exception:
            pass
        return ""

    @staticmethod
    def _notes(slide) -> str:  # type: ignore[no-untyped-def]
        """提取备注。没有备注时返回空串（不抛异常）。"""
        try:
            if not slide.has_notes_slide:
                return ""
            frame = slide.notes_slide.notes_text_frame
            return frame.text.strip() if frame is not None else ""
        except Exception:
            return ""

    def _table_block(self, table, index: int, slide_no: int) -> Block | None:  # type: ignore[no-untyped-def]
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)

        if not rows:
            return None

        width = max(len(r) for r in rows)
        norm = [r + [""] * (width - len(r)) for r in rows]
        header = norm[0]
        flat = [h if h else f"列{i + 1}" for i, h in enumerate(header)]

        cells: list[TableCell] = []
        for c, name in enumerate(flat):
            cells.append(TableCell(text=name, row=0, col=c, is_header=True))
        for r, row in enumerate(norm[1:], start=1):
            for c, val in enumerate(row):
                if val:
                    cells.append(TableCell(text=val, row=r, col=c))

        lines = [" | ".join(flat)]
        for row in norm[1:]:
            lines.append(" | ".join(row))

        return Block(
            type=BlockType.TABLE,
            text="\n".join(lines),
            page=slide_no,
            table_index=index,
            cells=cells,
            header_rows=1,
            extra={"slide": slide_no, "row_count": len(norm) - 1, "col_count": width},
        )
