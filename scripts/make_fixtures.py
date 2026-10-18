#!/usr/bin/env python3
"""生成 6 份"刁钻"合成样例文件，并同步产出期望输出 JSON。

## 为什么需要合成文件

公开文档覆盖"正常情况"，但解析器真正会翻车的是边界：
合并单元格、双行表头、纯图片 PDF、备注比正文多的 PPT、GBK 编码、空文件。

**这几种情况的共同点是：解析错了不会报错，而是静默给出错误结果。**
比"崩溃"危险得多。

## 设计原则

1. **内容完全虚构**：公司名、人名、数字都是编的，避免任何真实主体
2. **每个文件只测一个要点**：混在一起测，出问题不知道是哪里的锅
3. **期望输出就近生成**：造文件时同时写 expected/{name}.json，
   保证两者永远一致（手写期望值最容易和文件脱节）

用法：
    python scripts/make_fixtures.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "eval" / "fixtures" / "synthetic"
EXPECTED = ROOT / "eval" / "expected"

# 虚构主体，避免与任何真实公司/个人相关
FAKE_COMPANY = "青岚智造科技有限公司"
FAKE_PEOPLE = ["林绍安", "周慕云", "郑书言"]


# ---------------------------------------------------------------------------
# 1. 合并单元格 Excel
# ---------------------------------------------------------------------------
def make_xlsx_merged_cells() -> dict:
    """合并单元格 + 同一逻辑表的重复表头。

    测点：
    - 合并单元格只在左上角有值，其他是 None → 解析后每行都要带上分组名
    - 行级切块时，每行必须重复前置表头，否则"华东"那行的数字脱离上下文
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "区域销售"

    # 双行表头：第一行是分组，第二行是字段
    ws.append(["区域", "负责人", "季度销售", None, "备注"])
    ws.append([None, None, "Q1", "Q2", None])
    # 合并表头单元格
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:D1")
    ws.merge_cells("E1:E2")

    rows = [
        ("华东", "林绍安", 1280.5, 1432.8, "含税"),
        ("华东", "周慕云", 960.0, 1045.2, None),      # 同区域第二行，合并单元格值为空
        ("华南", "郑书言", 875.3, 812.6, None),
        ("华北", "林绍安", 1104.7, 1219.0, "含税"),   # 同名负责人在不同区域出现
    ]
    for r in rows:
        ws.append(list(r))

    # 合并"区域"列的同组单元格
    ws.merge_cells("A3:A4")

    out = FIXTURES / "xlsx_merged_cells.xlsx"
    wb.save(out)

    return {
        "file": out.name,
        "format": "xlsx",
        "test_focus": "合并单元格 + 双行表头 + 行级切块",
        "sheet_names": ["区域销售"],
        "expected_table": {
            "header_rows": [
                ["区域", "负责人", "季度销售", None, "备注"],
                [None, None, "Q1", "Q2", None],
            ],
            "data_row_count": 4,
            "first_data_row": ["华东", "林绍安", 1280.5, 1432.8, "含税"],
            # 关键断言：合并单元格解析后，第 4 行（华南前一行）的区域应仍是"华东"
            "merged_cell_inherited": {
                "row_index": 4,
                "column": "区域",
                "expected_value": "华东",
                "why": "A3:A4 合并，第 4 行原值为 None，解析后应继承为华东",
            },
        },
        "assertions": [
            "第一个 chunk 的 heading_path 含 '区域销售'",
            "存在包含 '华东' 与 '1280.5' 的 chunk",
            "存在包含 '华南' 与 '875.3' 的 chunk",
            "每个数据行 chunk 都重复携带表头（Q1/Q2 字段名）",
        ],
    }


# ---------------------------------------------------------------------------
# 2. 空行分隔的多张表
# ---------------------------------------------------------------------------
def make_xlsx_multi_tables() -> dict:
    """一个工作表里用空行分隔的两张逻辑表。

    测点：能否识别"空行 = 表边界"，而不是把两张表混成一张。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"

    ws.append(["产品编号", "名称", "单价"])
    ws.append(["P-001", "智能网关", 1280.0])
    ws.append(["P-002", "边缘控制器", 2450.0])
    ws.append([])  # 空行 = 表边界
    ws.append([])  # 连续两个空行
    ws.append(["园区", "装机数", "投运年份"])
    ws.append(["苏州园区", 128, 2024])
    ws.append(["合肥园区", 96, 2025])

    out = FIXTURES / "xlsx_multi_tables.xlsx"
    wb.save(out)

    return {
        "file": out.name,
        "format": "xlsx",
        "test_focus": "空行分隔的多张逻辑表",
        "sheet_names": ["汇总"],
        "expected_table": {
            "logical_table_count": 2,
            "table_1": {
                "header": ["产品编号", "名称", "单价"],
                "rows": 2,
                "first_row": ["P-001", "智能网关", 1280.0],
            },
            "table_2": {
                "header": ["园区", "装机数", "投运年份"],
                "rows": 2,
                "first_row": ["苏州园区", 128, 2024],
            },
        },
        "assertions": [
            "解析结果含两张表的表头，且不被合并成一张",
            "存在包含 'P-001' 的 chunk",
            "存在包含 '苏州园区' 的 chunk",
            "表 2 的 chunk 不应混入 'P-001'",
        ],
    }


# ---------------------------------------------------------------------------
# 3. 扫描件 PDF（纯图片，无文字层）
# ---------------------------------------------------------------------------
def make_pdf_scanned() -> dict:
    """模拟扫描件：页面里只有一张图片，没有文字层。

    测点：
    - 解析器必须检出"每页可提取字符 < 阈值"并明确报 no_text_layer
    - **绝不能返回空字符串然后当成功**

    实现方式：用 Pillow 画一张"扫描件"图，再用 pypdf 转成 PDF 页面。
    不手拼 PDF 字节 —— 那样容易生成非法文件（缺 xref / startxref），
    到时候测的是"文件坏了"而不是"解析器识别扫描件"。
    """
    import io

    from PIL import Image, ImageDraw
    from pypdf import PdfReader, PdfWriter

    # 画一张模拟扫描件的图：白底 + 灰色横条（模拟文字行），完全是图不是文字
    img = Image.new("RGB", (1240, 1754), "white")  # A4 @150dpi
    draw = ImageDraw.Draw(img)
    for i in range(18):
        y = 180 + i * 70
        draw.rectangle([120, y, 120 + 900 - (i % 5) * 90, y + 22], fill=(90, 90, 90))
    draw.rectangle([120, 140, 520, 175], fill=(40, 40, 40))  # 标题块

    img_buf = io.BytesIO()
    img.save(img_buf, format="PDF", resolution=150.0)
    img_buf.seek(0)

    # 用 pypdf 读入再写出，确保是合法 PDF
    reader = PdfReader(img_buf)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)

    out = FIXTURES / "pdf_scanned_no_text.pdf"
    with out.open("wb") as f:
        writer.write(f)

    # 自检：确保真的没有文字层，否则测试目标就落空了
    check = PdfReader(str(out))
    extracted = (check.pages[0].extract_text() or "").strip()
    if extracted:
        raise RuntimeError(
            f"生成的 PDF 意外含有文字层（{len(extracted)} 字符），无法用于测扫描件"
        )

    return {
        "file": out.name,
        "format": "pdf",
        "test_focus": "扫描件（无文字层）必须明确报错",
        "expected": {
            "should_raise": True,
            "error_code": "no_text_layer",
            "page_count": 1,
            "extractable_chars_per_page": 0,
        },
        "assertions": [
            "解析抛出 ParseError，code == 'no_text_layer'",
            "错误信息提示不支持 OCR",
            "不得返回空 ParsedDoc 并标记成功",
        ],
    }


# ---------------------------------------------------------------------------
# 4. PPTX：备注比正文多
# ---------------------------------------------------------------------------
def make_pptx_notes_heavy() -> dict:
    """每页正文一行，备注写了三倍的内容。

    测点：notes_slide 必须被提取。否则页面上"Q3 目标"这种短句
    与备注里的完整说明割裂，检索时只能命中无信息的标题。
    """
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slides_data = [
        {
            "title": "经营目标",
            "body": "2026 年 Q3 目标",
            "notes": (
                "本页目标是营收 8600 万元，毛利率不低于 32%。"
                "分产品看：智能网关贡献 5200 万，边缘控制器 2400 万，"
                "软件与运维服务 1000 万。"
            ),
        },
        {
            "title": "关键动作",
            "body": "三条主线",
            "notes": (
                "第一条是渠道下沉，新增 12 个地级市代理，重点覆盖苏北与皖北。"
                "第二条是产品迭代，9 月底前发布网关 3.2 版本，支持双链路冗余。"
                "第三条是交付提效，把现场实施周期从 21 天压到 14 天。"
            ),
        },
        {
            "title": "风险提示",
            "body": "两项风险",
            "notes": (
                "供应链风险：核心芯片交期从 8 周延长到 14 周，已锁定三季度用量。"
                "人力风险：交付工程师缺口 6 人，计划通过校招补充 4 人、外包补充 2 人。"
            ),
        },
    ]

    for item in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        slide.placeholders[1].text = item["body"]
        slide.notes_slide.notes_text_frame.text = item["notes"]

    out = FIXTURES / "pptx_notes_heavy.pptx"
    prs.save(out)

    return {
        "file": out.name,
        "format": "pptx",
        "test_focus": "notes_slide 提取（备注内容远多于正文）",
        "slide_count": 3,
        "expected": {
            "pages": [
                {"index": 1, "title": "经营目标", "notes_contains": ["8600 万元", "32%"]},
                {"index": 2, "title": "关键动作", "notes_contains": ["12 个地级市", "14 天"]},
                {"index": 3, "title": "风险提示", "notes_contains": ["14 周", "6 人"]},
            ]
        },
        "assertions": [
            "每页 chunk 都包含 notes 中的关键数字",
            "chunk 的 page 字段等于 slide 序号（1/2/3）",
            "备注内容不得丢失（含 8600 万元、12 个地级市、14 周）",
        ],
    }


# ---------------------------------------------------------------------------
# 5. GBK 编码 TXT
# ---------------------------------------------------------------------------
def make_txt_gbk() -> dict:
    """GBK 编码的中文文本，带 Windows 风格换行。

    测点：charset-normalizer 能否识别为 GBK。
    识别错了会得到一堆乱码，且**不会报错**。
    """
    lines = [
        "青岚智造科技有限公司 内部管理制度摘编",
        "",
        "一、差旅管理",
        "1. 住宿标准：一线城市每晚不超过 480 元，二线城市每晚不超过 350 元。",
        "2. 交通标准：市内交通凭票实报实销，单次不超过 120 元。",
        "3. 出差补贴：每人每天 120 元，按实际出差天数计算。",
        "",
        "二、采购管理",
        "1. 单笔采购金额超过 5 万元需经分管副总审批。",
        "2. 单笔采购金额超过 20 万元需经总经理审批并附三家比价。",
        "",
        "三、用印管理",
        "1. 合同用印需在 OA 提交申请，审批通过后由行政部统一加盖。",
        "2. 空白文件一律不得用印。",
    ]
    content = "\r\n".join(lines)

    out = FIXTURES / "txt_gbk_chinese.txt"
    out.write_bytes(content.encode("gbk"))

    return {
        "file": out.name,
        "format": "txt",
        "test_focus": "GBK 编码识别",
        "expected": {
            "encoding": "gbk",
            "line_count": len(lines),
            "sections": ["一、差旅管理", "二、采购管理", "三、用印管理"],
            "key_facts": {
                "一线城市住宿上限": "480 元",
                "出差补贴": "120 元",
                "采购审批阈值_副总": "5 万元",
                "采购审批阈值_总经理": "20 万元",
            },
        },
        "assertions": [
            "解析出的第一个字符是 '青'，不是乱码",
            "文本包含 '480 元' 与 '20 万元'",
            "不得因编码错误产生替换字符 U+FFFD",
        ],
    }


# ---------------------------------------------------------------------------
# 6. 空文件
# ---------------------------------------------------------------------------
def make_empty_file() -> dict:
    """0 字节文件。

    测点：必须明确拒绝，而不是塞一个空 chunk 进索引。
    空 chunk 会污染 BM25 的文档长度统计。
    """
    out = FIXTURES / "empty.txt"
    out.write_bytes(b"")

    return {
        "file": out.name,
        "format": "txt",
        "test_focus": "空文件必须拒绝",
        "expected": {
            "should_raise": True,
            "error_code": "empty_file",
            "size_bytes": 0,
        },
        "assertions": [
            "解析抛出 ParseError，code == 'empty_file'",
            "不得生成任何 chunk",
        ],
    }


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------
MAKERS = [
    make_xlsx_merged_cells,
    make_xlsx_multi_tables,
    make_pdf_scanned,
    make_pptx_notes_heavy,
    make_txt_gbk,
    make_empty_file,
]


def main() -> int:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    EXPECTED.mkdir(parents=True, exist_ok=True)

    print(f"输出目录：\n  {FIXTURES}\n  {EXPECTED}\n")

    index = []
    for maker in MAKERS:
        try:
            meta = maker()
            name = Path(meta["file"]).stem
            (EXPECTED / f"{name}.json").write_text(
                json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            size = (FIXTURES / meta["file"]).stat().st_size
            print(f"  ✓ {meta['file']:<28} {size:>8} bytes  — {meta['test_focus']}")
            index.append(
                {
                    "file": meta["file"],
                    "format": meta["format"],
                    "test_focus": meta["test_focus"],
                }
            )
        except Exception as e:
            print(f"  ✗ {maker.__name__} 失败：{type(e).__name__}: {e}", file=sys.stderr)
            return 1

    (EXPECTED / "_index.json").write_text(
        json.dumps(
            {
                "generated_by": "scripts/make_fixtures.py",
                "note": "合成样例，内容全部虚构；期望输出与文件同源生成，不要手改其中一个",
                "items": index,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"\n共 {len(index)} 份，期望输出已写入 {EXPECTED}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
